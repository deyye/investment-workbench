"""Deterministic comparison and formula-safe Excel export."""
import io,re
from decimal import Decimal
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from .extract import FIELDS,STAGES,METRICS,clean,numeric,MISSING_REASONS
from .quantities import quantity_key,rounding_equivalent
from datetime import datetime

COLORS={'same':'FFFFFF','different':'FFF0D7','equivalent':'E7F1FF','missing':'EEF1F5','review':'FCE7E8',
        'align':'EDE7F6','unextracted':'FDE7E7'}
LABELS={'same':'一致','different':'内容变化','equivalent':'表述不同 / 标准值相同','missing':'未载明',
        'review':'待核对','align':'疑似同一指标','unextracted':'原文有线索未提取到'}

def norm(value):return clean(value).replace('〔','[').replace('〕',']').replace('，',',').replace('。','')
def compare_cells(cells):
    values=[c.get('value') for c in cells]
    present=[v for v in values if v is not None and v!='']
    if any(c.get('status') in ['needs_review','conflict','uncertain'] for c in cells):status='review'
    elif len(present)!=len(values) and len(set(present))<=1:status='missing'
    elif len(set(values))<=1:status='same'
    elif len({norm(v) for v in present})==1:status='equivalent'
    else:status='different'
    # 「未载明」要分两种：原文确无该要素（absent）与原文有线索却没抽到（unextracted）。
    # 混在一起会让正确的缺失看起来像系统故障，也让真正该处理的漏抽淹没在灰格里。
    if status=='missing' and any(c.get('reason')=='unextracted' for c in cells):status='unextracted'
    return status

def date_key(value):
    m=re.search(r'(\d{4})年(\d{1,2})月(\d{1,2})日',value or '')
    return tuple(map(int,m.groups())) if m else (0,0,0)

def rows_for(docs):
    # 兜底下标取 len(STAGES)-1（=「待确认」），不写死数字：STAGES 增删阶段时
    # 写死的 3 会悄悄变成别的阶段的位次，把未知阶段排到中间去。
    docs=sorted(docs,key=lambda d:(STAGES.index(d['stage']) if d['stage'] in STAGES else len(STAGES)-1,date_key(d['fields']['印发日期'].get('value')),d['filename']))
    rows=[]
    for name in FIELDS:
        cs=[d['fields'][name] for d in docs]
        status=compare_cells(cs);note=''
        if name in ['建设周期','总投资/匡算/估算/概算'] and status not in ['review','missing','unextracted']:
            ns=[numeric(c.get('value')) for c in cs]
            bases=[re.search(r'匡算|估算|概算',c.get('value') or '') for c in cs]
            basis=[m.group() if m else '总投资' for m in bases]
            if all(ns) and len({quantity_key(n) for n in ns})==1:
                if name=='建设周期' or len(set(basis))==1:
                    status='equivalent' if len({c['value'] for c in cs})>1 else status
            # 与指标同一套书写精度宽容：金额/工期只因四舍五入而略差时，
            # 不报成「内容变化」，但仍保留原有投资口径提示。
            if status=='different' and (name=='建设周期' or len(set(basis))==1) and rounding_equivalent(ns,cs):
                status='equivalent'
                note=_rounding_note(ns,[c.get('value') for c in cs])
            if name=='总投资/匡算/估算/概算' and len(set(basis))>1:note='投资口径：'+' → '.join(basis)+'；阶段金额变化不直接认定为异常。'
        rows.append({'name':name,'kind':'fixed','cells':cs,'status':status,'note':note})
    names=list(dict.fromkeys(m['name'] for d in docs for m in d['metrics']))
    priority={name:i for i,(name,_) in enumerate(METRICS)}
    names.sort(key=lambda name:priority.get(name,len(priority)))
    for name in names:
        cs=[]
        for d in docs:
            hits=[m for m in d['metrics'] if m['name']==name]
            if not hits:cs.append({'value':None,'status':'missing','evidence':[]})
            elif len(hits)==1:cs.append(hits[0])
            else:cs.append({'value':'；'.join(m['value'] for m in hits),'status':'conflict' if len({m['value'] for m in hits})>1 else ('reviewed' if all(m.get('status')=='reviewed' for m in hits) else 'needs_review'),'evidence':[e for m in hits for e in m['evidence']]})
        status=compare_cells(cs);note=''
        ns=[numeric(c['value']) if c.get('value') else None for c in cs]
        available=[(i,n) for i,n in enumerate(ns) if n is not None]
        if len(available)>1 and status not in ('review','unextracted'):
            indices=[i for i,n in available]
            vals=[n for i,n in available]
            keys=[quantity_key(n) for n in vals]
            if len(set(keys))==1 and len({c['value'] for c in cs})>1:status='equivalent' if len(available)==len(cs) else 'missing'
            # 书写精度宽容（须排在精确等价判断之后）：真实批复里同一块地写成
            # 「13.693公顷」与「136929.6平方米」只差 0.4 平方米，若报成
            # 「内容变化」会逼人工逐条判断，属噪音而非信息。
            elif rounding_equivalent(vals,[cs[i] for i in indices]):
                status='equivalent' if len(available)==len(cs) else 'missing'
                note=_rounding_note(vals,[cs[i].get('value') for i in indices])
            if not note and len({n['unit'] for n in vals})==1 and not any(n['upper'] for n in vals):
                notes=[]
                for i in range(1,len(vals)):
                    a,b=Decimal(vals[i-1]['number']),Decimal(vals[i]['number'])
                    if a==b:continue
                    step=f'{docs[indices[i-1]]["stage"]}→{docs[indices[i]]["stage"]}：'
                    # 单步差异若只是书写精度/约数造成，就不要给出百分比——
                    # 实测里「-0.400平方米（变动比例小于0.01%）」这种提示，
                    # 用户只能逐条点开才能确认它无需处理，是纯粹的噪音。
                    if rounding_equivalent([vals[i-1],vals[i]],[cs[indices[i-1]],cs[indices[i]]]):
                        notes.append(step+f'{b-a:+f}{vals[i]["unit"]}（差异属书写精度或约数范围，未认定为变化）')
                        continue
                    note_i=step+f'{b-a:+f}{vals[i]["unit"]}'
                    if a:
                        pct=(b-a)/a*100
                        note_i+=('（变动比例小于0.01%）' if abs(pct)<Decimal('.005') else f'（{pct:+.2f}%）')
                    if vals[i-1]['qualifier'] or vals[i]['qualifier']:note_i+='，含约数/限定'
                    notes.append(note_i)
                note='；'.join(notes)
        rows.append({'name':name,'kind':'metric','cells':cs,'status':status,'note':note})
    decided={h['name'] for d in docs for h in d.get('history',[]) if h.get('kind')=='align'}
    _mark_alignment(rows,decided)
    return docs,rows

def _rounding_note(items,raw_values=None):
    """精度/约数差异的说明文案：把「差多少」「各阶段怎么写」和「为什么不算变化」都讲清楚。"""
    nums=[Decimal(n['number']) for n in items]
    spread=max(nums)-min(nums)
    shown=' 与 '.join(dict.fromkeys(v for v in (raw_values or []) if v))
    if shown:
        return ('各阶段写法精度不同（%s），极差 %s%s 属书写精度或约数范围，未认定为内容变化。'
                % (shown,'%.10g'%spread,items[0]['unit']))
    return ('极差 %s%s 属书写精度或约数范围，未认定为内容变化。' % ('%.10g'%spread,items[0]['unit']))

def _metric_key(cell):
    """把一行的可比值收敛成单个比较键；一行内出现多个不同值（冲突）时返回 None。"""
    vals=[quantity_key(numeric(c['value'])) for c in cell if c.get('value') and numeric(c['value'])]
    if not vals:return None
    return vals[0] if len(set(vals))==1 else None

def _mark_alignment(rows,decided=None):
    """标出「数值相同、只是名称不同、且各出现在不同文档列」的指标行。

    实测动机：龙泉项目可行性研究写「道路路面面积60637平方米」，初步设计写
    「行车道、路缘带及硬路肩总面积60637平方米」——同一数值、同一工程，只因
    两个阶段的措辞不同就被拆成两行，各自显示「未载明」。看表的人会得出
    「这个阶段没有该指标」的错误结论，而实际两个阶段都写了。

    只在两行**文档列互不重叠**时才提示：同一份文档里的「涵洞7处」与「平交口7处」
    数值相同，那不是跨阶段错位，不应提示。命中只标「疑似」，是否需要合并交人工决定。
    """
    metrics=[r for r in rows if r['kind']=='metric']
    shaped=[]
    for r in metrics:
        idx=frozenset(i for i,c in enumerate(r['cells']) if c.get('value'))
        if not idx:continue
        key=_metric_key(r['cells'])
        if key is None:continue
        shaped.append((r,idx,key))
    for i in range(len(shaped)):
        for j in range(i+1,len(shaped)):
            ra,ia,ka=shaped[i];rb,ib,kb=shaped[j]
            if ka!=kb or ra['name']==rb['name']:continue
            if '｜'.join(sorted((ra['name'],rb['name']))) in (decided or set()):continue
            if ia & ib:continue
            ra.setdefault('align_with',set()).add(rb['name'])
            rb.setdefault('align_with',set()).add(ra['name'])
    for r in metrics:
        if not r.get('align_with'):continue
        others=sorted(r['align_with'])
        r['align_with']=others
        value=next((c['value'] for c in r['cells'] if c.get('value')),'')
        r['status']='align'
        r['note']=('与「%s」数值相同（%s），疑似同一指标，建议人工确认是否合并。'
                   '当前分列显示会让两行都看起来像「未载明」。' % ('」「'.join(others),value))

def safe(value):
    s='' if value is None else str(value)
    return "'"+s if s.lstrip().startswith(('=','+','-','@')) else s

def export_xlsx(documents):
    wb=Workbook();wb.remove(wb.active)
    groups={}
    for d in documents:groups.setdefault(d['project_key'],[]).append(d)
    evidence=wb.create_sheet('证据索引');evidence.append(['项目','文件','阶段','字段','值','PDF页码','原文','坐标'])
    # 「来源」列是这套系统可审计的基础：规则抽的、人工改的、大模型给的、
    # 还是自检循环补救回来的，导出后必须能分辨，否则"每个值都要有出处"落不了地。
    single=wb.create_sheet('单文档结构化');single.append(['项目','文件','阶段','字段','提取值','来源','状态','缺失原因','证据页码'])
    history=wb.create_sheet('修订记录');history.append(['文件','时间','字段','修改前','修改后','说明'])
    review=wb.create_sheet('待复核事项');review.append(['项目','文件','事项'])
    for gi,ds in enumerate(groups.values(),1):
        docs,rows=rows_for(ds)
        for kind,label in [('fixed','固定字段'),('metric','建设指标')]:
            ws=wb.create_sheet(f'{gi}-{label}')
            ws.append([safe(docs[0]['project_name'])]);ws.append(['字段']+[d['stage']+' | '+d['filename'] for d in docs]+['比较结果','说明'])
            for r in rows:
                if r['kind']!=kind:continue
                ws.append([safe(r['name'])]+[safe(c.get('value') if c.get('value') is not None else '未载明') for c in r['cells']]+[LABELS[r['status']],safe(r['note'])])
                for c in ws[ws.max_row]:c.fill=PatternFill('solid',fgColor=COLORS['missing' if c.value=='未载明' else r['status']])
            ws.freeze_panes='B3';ws.auto_filter.ref=f'A2:{ws.cell(ws.max_row,ws.max_column).coordinate}'
        for d in docs:
            for k,c in list(d['fields'].items())+[(m['name'],m) for m in d['metrics']]:
                single.append([safe(d['project_name']),safe(d['filename']),d['stage'],safe(k),safe(c.get('value')),
                               safe(c.get('method')),c.get('status'),MISSING_REASONS.get(c.get('reason'),''),
                               ','.join(str(p) for p in sorted({e['page'] for e in c['evidence']}))])
                for e in c['evidence']:
                    evidence.append([safe(d['project_name']),safe(d['filename']),d['stage'],safe(k),safe(c.get('value')),e['page'],safe(e['quote']),str(e['bbox'])])
                if c.get('status') in ['needs_review','conflict','uncertain']:
                    review.append([safe(d['project_name']),safe(d['filename']),safe(k+'：待核对')])
                elif c.get('status')=='missing' and c.get('reason')=='unextracted':
                    review.append([safe(d['project_name']),safe(d['filename']),safe(k+'：原文有线索但未提取到，请人工确认')])
            for r in rows:
                if r['kind']=='metric' and r['status']=='align':
                    review.append([safe(d['project_name']),safe(d['filename']),safe(r['name']+'：'+r['note'])])
            for h in d.get('history',[]):
                history.append([safe(d['filename']),datetime.fromtimestamp(h['time']).isoformat(),safe(h.get('name','')),safe(h.get('before',{}).get('value') if isinstance(h.get('before'),dict) else h.get('before')),safe(h.get('after')),safe(h.get('reason',''))])
            for w in d['warnings']:review.append([safe(d['project_name']),safe(d['filename']),safe(w)])
    for ws in wb:
        for row in ws:
            for c in row:c.alignment=Alignment(vertical='top',wrap_text=True)
        head=2 if '-' in ws.title else 1
        for c in ws[head]:c.fill=PatternFill('solid',fgColor='173A60');c.font=Font(color='FFFFFF',bold=True)
        for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=38 if col[0].column>1 else 25
        if ws.title=='证据索引':ws.column_dimensions['G'].width=75
    out=io.BytesIO();wb.save(out);return out.getvalue()
