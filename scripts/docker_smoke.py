"""HTTP acceptance against a running container. Creates only a synthetic fixture."""
import argparse
import base64
import io
import json
import time
import urllib.request
from pathlib import Path
import pymupdf as fitz
from openpyxl import load_workbook

parser=argparse.ArgumentParser()
parser.add_argument('--url',default='http://127.0.0.1:8765')
parser.add_argument('--state',default='/tmp/approval-smoke-id.json')
parser.add_argument('--verify-persistence',action='store_true')
a=parser.parse_args()

def call(path,data=None):
    request=urllib.request.Request(a.url+path,data=json.dumps(data).encode() if data is not None else None,headers={'Content-Type':'application/json','X-Requested-With':'ApprovalAgent'})
    with urllib.request.urlopen(request,timeout=30) as r:return r.read()

assert json.loads(call('/api/health'))['status']=='ok'
assert '批文研析' in call('/').decode()
if a.verify_persistence:
    saved=json.loads(Path(a.state).read_text());docid=saved['id']
    restored=json.loads(call('/api/jobs/'+saved['job_id']))
    assert restored['status']=='completed'
    assert restored['files'][0]['status']=='success'
else:
    with fitz.open() as pdf:
        page=pdf.new_page();page.insert_text((50,70),'Synthetic Docker acceptance document')
        blob=pdf.tobytes()
    job=json.loads(call('/api/upload',{'files':[{'name':'docker-fixture.pdf','data':base64.b64encode(blob).decode()}]}))
    deadline=time.monotonic()+60
    while True:
        result=json.loads(call('/api/jobs/'+job['job_id']))
        if result['status']=='completed':break
        if time.monotonic()>deadline:raise RuntimeError('Container processing timeout')
        time.sleep(.25)
    assert not result['errors'],result['errors']
    docid=result['results'][0]['id']
    call('/api/documents/'+docid+'/review',{'kind':'fixed','name':'项目名称','value':'虚构容器验收项目'})
    Path(a.state).write_text(json.dumps({'id':docid,'job_id':job['job_id']}))
assert json.loads(call('/api/documents/'+docid))['fields']['项目名称']['value']=='虚构容器验收项目'
assert call('/api/documents/'+docid+'/pages/1.png').startswith(b'\x89PNG')
workbook=load_workbook(io.BytesIO(call('/api/export.xlsx')))
assert '证据索引' in workbook.sheetnames
print(json.dumps({'ok':True,'persistence':a.verify_persistence,'checks':['health','UI','stored review','PDF render','Excel']}))
