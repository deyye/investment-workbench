import io,json,os,tempfile,unittest
from pathlib import Path
from unittest.mock import patch
import fitz
from openpyxl import load_workbook
from app.extract import numeric,augment_llm,extract,FIELDS
from app.compare import compare_cells,rows_for,export_xlsx,safe

def c(value,status='extracted'):return {'value':value,'status':status,'evidence':[]}

class CoreTests(unittest.TestCase):
    def test_area_conversion(self):
        self.assertEqual(numeric('13.693公顷')['number'],'136930.000')
        self.assertEqual(numeric('13.693公顷')['unit'],'平方米')
    def test_speed_unit_not_length(self):
        self.assertEqual(numeric('60公里/小时')['unit'],'千米/小时')
        self.assertEqual(numeric('60公里/小时')['number'],'60')
    def test_range_with_repeated_units(self):
        self.assertEqual(numeric('24.5米-36米')['upper'],'36')
        self.assertEqual(numeric('24.5-36米')['upper'],'36')
    def test_approximate_retained(self):self.assertEqual(numeric('约2.31公里')['qualifier'],'约')
    def test_missing_is_not_same(self):self.assertEqual(compare_cells([c('14个月'),c(None,'missing')]),'missing')
    def test_missing_does_not_hide_other_changes(self):self.assertEqual(compare_cells([c('3'),c(None,'missing'),c('4')]),'different')
    def test_formatting(self):self.assertEqual(compare_cells([c('浙〔2025〕1号'),c('浙[2025]1号')]),'equivalent')
    def test_review_priority(self):self.assertEqual(compare_cells([c('有','needs_review'),c('有')]),'review')
    def test_excel_injection(self):
        for s in ['=1+1',' @SUM(A1)','+1','-1']:self.assertTrue(safe(s).startswith("'"))
    def test_llm_requires_real_evidence(self):
        r={'lines':[{'id':'p1-l1','text':'总建筑面积3994平方米','page':1,'bbox':[1,2,3,4]}], 'fields':{k:c(None) for k in FIELDS},'metrics':[],'warnings':[]}
        out={'fields':{'项目名称':{'value':'虚构项目','evidence_ids':['p1-l1']}},'metrics':[{'name':'总建筑面积','value':'3994平方米','evidence_ids':['p1-l1']},{'name':'虚构面积','value':'888平方米','evidence_ids':['p9-l9']}]}
        with patch('app.extract.chat',return_value=out):augment_llm(r)
        self.assertIsNone(r['fields']['项目名称']['value']);self.assertEqual(len(r['metrics']),1)
    def test_no_cross_project_merge(self):
        docs=[]
        for i in range(2):docs.append({'id':str(i),'stage':'初步设计','filename':str(i),'project_key':str(i),'project_name':'同名项目','fields':{k:c(None) for k in FIELDS},'metrics':[],'warnings':[]})
        wb=load_workbook(io.BytesIO(export_xlsx(docs)))
        self.assertIn('1-固定字段',wb.sheetnames);self.assertIn('2-固定字段',wb.sheetnames)
    def test_scanned_or_empty_not_hallucinated(self):
        with tempfile.TemporaryDirectory() as td:
            p=Path(td)/'blank.pdf';d=fitz.open();d.new_page();d.save(p);d.close()
            r=extract(p,p.name,'0'*32)
        self.assertIsNone(r['fields']['总投资/匡算/估算/概算']['value']);self.assertEqual(r['stage'],'待确认');self.assertEqual(r['fields']['印章']['value'],'无法判断')

@unittest.skipUnless(os.getenv('SAMPLE_DIR'),'Set SAMPLE_DIR to run six supplied PDF regression tests')
class SuppliedPDFTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.docs={p.name:extract(p,p.name,f'{i:032x}') for i,p in enumerate(Path(os.environ['SAMPLE_DIR']).glob('*.pdf'))}
    def test_document_numbers(self):
        expected={'初步设计批复文件.pdf':'庆发改投〔2026〕1号','可行性研究报告批复文件.pdf':'庆发改投[2025]148号','项目建议书批复文件.pdf':'庆发改投〔2025〕127号','初步设计批复文件(1).pdf':'龙发改投资〔2023〕192号','可行性研究报告批复文件(1).pdf':'龙发改投资〔2023〕183号','项目建议书批复文件(1).pdf':'龙发改投资〔2022〕182号'}
        for fn,v in expected.items():
            with self.subTest(fn=fn):self.assertEqual(self.docs[fn]['fields']['发文字号']['value'],v)
    def test_all_14_fields_and_evidence(self):
        self.assertEqual(len(self.docs),6)
        for d in self.docs.values():
            self.assertEqual(set(d['fields']),set(FIELDS))
            for k,c in d['fields'].items():
                if c['value'] is not None:self.assertTrue(c['evidence'],(d['filename'],k))
                for e in c['evidence']:
                    p=d['pages'][e['page']-1];x0,y0,x1,y1=e['bbox']
                    self.assertTrue(0<=x0<x1<=p['width']+1);self.assertTrue(0<=y0<y1<=p['height']+1)
    def test_seals(self):
        for d in self.docs.values():self.assertEqual(d['fields']['印章']['value'],'有')
    def test_qingyuan_area_sequence(self):
        vals=[]
        for fn in ['项目建议书批复文件.pdf','可行性研究报告批复文件.pdf','初步设计批复文件.pdf']:
            vals.append(next(m['value'] for m in self.docs[fn]['metrics'] if m['name']=='总建筑面积'))
        self.assertEqual(vals,['约4702平方米','3994平方米','3576.63平方米'])
    def test_no_imprint_fabrication(self):
        for fn,d in self.docs.items():
            expected='龙泉市发展和改革局办公室' if '(1)' in fn else None
            self.assertEqual(d['fields']['印发机关']['value'],expected)
    def test_missing_duration(self):
        self.assertIsNone(self.docs['初步设计批复文件.pdf']['fields']['建设周期']['value'])
        self.assertIsNone(self.docs['项目建议书批复文件.pdf']['fields']['建设周期']['value'])
        self.assertEqual(self.docs['可行性研究报告批复文件.pdf']['fields']['建设周期']['value'],'14个月')
    def test_longquan_ranges(self):
        d=self.docs['初步设计批复文件(1).pdf']
        m=next(m for m in d['metrics'] if m['name']=='路基宽度')
        self.assertEqual(m['normalized']['upper'],'36')
    def test_no_page_number_noise(self):
        for d in self.docs.values():self.assertNotIn('——',d['fields']['建设内容']['value'])
    def test_excel_values_and_colours(self):
        wb=load_workbook(io.BytesIO(export_xlsx(list(self.docs.values()))))
        sheets=[ws for ws in wb if ws.title.endswith('建设指标')]
        area=[row for ws in sheets for row in ws.iter_rows() if row[0].value=='总建筑面积'][0]
        self.assertEqual(area[1].value,'约4702平方米');self.assertEqual(area[1].fill.fgColor.rgb,'00FFF0D7')
    def test_evidence_quotes_exist(self):
        for d in self.docs.values():
            lines={l['id']:l for l in d['lines']}
            for c in list(d['fields'].values())+d['metrics']:
                for e in c['evidence']:
                    if '-seal-' not in e['id']:self.assertIn(e['quote'],lines[e['id']]['text'])

if __name__=='__main__':unittest.main()
