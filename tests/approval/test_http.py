"""Real HTTP contract, isolated temporary storage, no external services."""
import base64,io,json,tempfile,threading,time,unittest,urllib.request,urllib.error
from http.server import ThreadingHTTPServer
import pymupdf as fitz
from openpyxl import load_workbook
from app.server import Handler,Store

class HTTPTests(unittest.TestCase):
    def test_upload_job_preview_review_export_and_duplicate(self):
        with tempfile.TemporaryDirectory() as td:
            server=ThreadingHTTPServer(('127.0.0.1',0),Handler);server.store=Store(td)
            t=threading.Thread(target=server.serve_forever,daemon=True);t.start()
            url=f'http://127.0.0.1:{server.server_port}'
            def call(path,data=None):
                req=urllib.request.Request(url+path,data=json.dumps(data).encode() if data else None,headers={'Content-Type':'application/json','X-Requested-With':'ApprovalAgent'})
                return urllib.request.urlopen(req,timeout=20)
            try:
                with call('/') as r:self.assertIn('批文研析',r.read().decode())
                # A generated independent fixture, not copied sample output.
                d=fitz.open();p=d.new_page();p.insert_text((70,80),'Independent fixture document for API validation.');blob=d.tobytes();d.close()
                payload={'files':[{'name':'fixture.pdf','data':base64.b64encode(blob).decode()}],'use_llm':False}
                def upload():
                    with call('/api/upload',payload) as r:jid=json.load(r)['job_id']
                    for _ in range(150):
                        with call('/api/jobs/'+jid) as r:j=json.load(r)
                        if j['status']=='completed':return j
                        time.sleep(.1)
                    self.fail('job timeout')
                j=upload();self.assertFalse(j['errors']);docid=j['results'][0]['id']
                with call('/api/documents/'+docid+'/pages/1.png') as r:self.assertTrue(r.read().startswith(b'\x89PNG'))
                with call('/api/documents/'+docid+'/review',{'kind':'fixed','name':'项目名称','value':'人工确认项目'}) as r:self.assertTrue(json.load(r)['ok'])
                with call('/api/documents') as r:gs=json.load(r)['groups']
                self.assertEqual(gs[0]['name'],'人工确认项目');self.assertEqual(len(gs[0]['documents'][0]['history']),1)
                with call('/api/export.xlsx') as r:wb=load_workbook(io.BytesIO(r.read()))
                self.assertIn('证据索引',wb.sheetnames)
                self.assertTrue(upload()['results'][0]['duplicate'])
                bad=urllib.request.Request(url+'/api/upload',data=json.dumps(payload).encode(),headers={'Content-Type':'application/json'})
                with self.assertRaises(urllib.error.HTTPError) as err:urllib.request.urlopen(bad)
                self.assertEqual(err.exception.code,403)
                with self.assertRaises(urllib.error.HTTPError):call('/api/upload',{'files':[{'name':'x.pdf','data':base64.b64encode(b'not pdf').decode()}]})
            finally:
                server.shutdown();server.server_close();server.store.executor.shutdown(wait=True);t.join()
